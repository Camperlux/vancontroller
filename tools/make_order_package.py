"""Build the PCBWay turnkey order package (fabrication + assembly + parts sourcing).

Usage:  python tools/make_order_package.py
Output: manufacturing/VanController_rev0.3_PCBWay/  and a zip of it.

Run tools/make_fab.py first (Gerbers, drill, PDFs). This script sets the board's drill/place origin to the
bottom-left board corner (via KiCad's Python) so the pick-and-place coordinates are board-relative.
"""
import csv
import os
import shutil
import subprocess
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KI = r"C:\Program Files\KiCad\10.0\bin"
CLI = os.path.join(KI, "kicad-cli.exe")
KIPY = os.path.join(KI, "python.exe")
PCB = os.path.join(PRJ, "VanController.kicad_pcb")
SCH = os.path.join(PRJ, "VanController.kicad_sch")
REV = "rev0.3"
OUT = os.path.join(PRJ, "manufacturing", f"VanController_{REV}_PCBWay")
BOARD_ORIGIN = (40.0, 175.0)  # page coordinates of the bottom-left board corner (board is 230 x 135 at 40,40)

# Through-hole parts are sourced by PCBWay but shipped loose; the customer solders them.
SELF_THT = True
THT_PKG = ("Relay_SPDT_Omron-G5LE-1", "PhoenixContact_MSTBA", "Fuseholder_Blade", "FuseHolder_Blade", "PinHeader_")

# Parts not placed by the assembler. Refs here are excluded from the pick-and-place file.
DNP = {
    "J1": "M5 stud: customer fits an M5 bolt + nuts (ring-lug terminal)",
    "J2": "M5 stud: customer fits an M5 bolt + nuts (ring-lug terminal)",
}

# (value, footprint) -> (manufacturer, MPN, description, mount)
# Passive MPNs are examples to the stated spec; equivalents to the same spec are acceptable.
PARTS = {
    ("100nF", "C_0603_1608Metric"): ("Samsung", "CL10B104KB8NNNC", "Cap MLCC 100nF 50V X7R 0603", "SMD"),
    ("100nF 100V", "C_0603_1608Metric"): ("Murata", "GRM188R72A104KA35D", "Cap MLCC 100nF 100V X7R 0603", "SMD"),
    ("100uF 35V", "CP_Elec_8x10"): ("Panasonic", "EEE-FK1V101P", "Cap Al-elec 100uF 35V 8x10.2mm SMD", "SMD"),
    ("4.7uF 50V", "C_1210_3225Metric"): ("Murata", "GRM32ER71H475KA88L", "Cap MLCC 4.7uF 50V X7R 1210", "SMD"),
    ("4.7uF 100V", "C_1210_3225Metric"): ("TDK", "C3225X7S2A475M200AB", "Cap MLCC 4.7uF 100V X7S 1210", "SMD"),
    ("4.7nF", "C_0603_1608Metric"): ("Samsung", "CL10B472KB8NNNC", "Cap MLCC 4.7nF 50V X7R 0603", "SMD"),
    ("47pF", "C_0603_1608Metric"): ("Samsung", "CL10C470JB8NNNC", "Cap MLCC 47pF 50V C0G 0603", "SMD"),
    ("47uF 10V", "C_1210_3225Metric"): ("Murata", "GRM32ER61A476KE20L", "Cap MLCC 47uF 10V X5R 1210", "SMD"),
    ("22uF", "C_0805_2012Metric"): ("Murata", "GRM21BR61A226ME44L", "Cap MLCC 22uF 10V X5R 0805", "SMD"),
    ("1uF", "C_0603_1608Metric"): ("Samsung", "CL10B105KA8NNNC", "Cap MLCC 1uF 25V X7R 0603", "SMD"),
    ("10nF", "C_0603_1608Metric"): ("Samsung", "CL10B103KB8NNNC", "Cap MLCC 10nF 50V X7R 0603", "SMD"),
    ("10uF", "C_0805_2012Metric"): ("Samsung", "CL21A106KAYNNNE", "Cap MLCC 10uF 25V X5R 0805", "SMD"),
    ("SMDJ26CA", "D_SMC"): ("Littelfuse", "SMDJ26CA", "TVS 3000W 26V bidirectional SMC", "SMD"),
    ("B360A", "D_SMA"): ("Diodes Inc", "B360A-13-F", "Schottky 60V 3A SMA", "SMD"),
    ("B5819W", "D_SOD-123"): ("Diodes Inc", "B5819W-7-F", "Schottky 40V 1A SOD-123", "SMD"),
    ("1N4148W", "D_SOD-123"): ("Diodes Inc", "1N4148W-7-F", "Switching diode 100V SOD-123", "SMD"),
    ("BAT54S", "SOT-23"): ("Nexperia", "BAT54S,215", "Schottky dual series 30V SOT-23", "SMD"),
    ("SS54", "D_SMC"): ("Vishay", "SS54-E3/57T", "Schottky 40V 5A SMC", "SMD"),
    ("NUP2105L", "SOT-23"): ("onsemi", "NUP2105LT1G", "CAN bus ESD protection SOT-23", "SMD"),
    ("SM712", "SOT-23"): ("Semtech", "SM712.TCT", "RS-485 TVS array SOT-23", "SMD"),
    ("Green", "LED_0603_1608Metric"): ("Lite-On", "LTST-C191KGKT", "LED green 0603", "SMD"),
    ("Blue", "LED_0603_1608Metric"): ("Lite-On", "LTST-C191TBKT", "LED blue 0603", "SMD"),
    ("Red", "LED_0603_1608Metric"): ("Lite-On", "LTST-C191KRKT", "LED red 0603", "SMD"),
    ("Yellow", "LED_0603_1608Metric"): ("Lite-On", "LTST-C191KSKT", "LED yellow 0603", "SMD"),
    ("30A ATO", "FuseHolder_Blade_ATO_Littelfuse_FLR_178.6165"):
        ("Littelfuse", "FLR 178.6165 series (verify full PN)", "PCB fuse holder, ATO blade, 30 A. Fuse NOT fitted", "THT"),
    ("10A mini", "Fuseholder_Blade_Mini_Keystone_3568"):
        ("Keystone", "3568", "PCB fuse holder, mini blade. Fuse NOT fitted", "THT"),
    ("5A mini", "Fuseholder_Blade_Mini_Keystone_3568"):
        ("Keystone", "3568", "PCB fuse holder, mini blade. Fuse NOT fitted", "THT"),
    ("USB-C", "USB_C_Receptacle_GCT_USB4105-xx-A_16P_TopMnt_Horizontal"):
        ("GCT", "USB4105-GF-A", "USB-C receptacle USB2.0 16P top mount", "SMD+THT"),
    ("UART0 debug", "PinHeader_1x04_P2.54mm_Vertical"): ("Any", "2.54mm 1x04 male header, gold", "Pin header 1x04 vertical", "THT"),
    ("AUX ADC", "PinHeader_1x04_P2.54mm_Vertical"): ("Any", "2.54mm 1x04 male header, gold", "Pin header 1x04 vertical", "THT"),
    ("NO / COM / NC", "PhoenixContact_MSTBA_2,5_3-G-5,08_1x03_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757255", "MSTBA 2,5/ 3-G-5,08 header 3-pos 5.08mm", "THT"),
    ("CANH / CANL / GND", "PhoenixContact_MSTBA_2,5_3-G-5,08_1x03_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757255", "MSTBA 2,5/ 3-G-5,08 header 3-pos 5.08mm", "THT"),
    ("A / B / GND", "PhoenixContact_MSTBA_2,5_3-G-5,08_1x03_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757255", "MSTBA 2,5/ 3-G-5,08 header 3-pos 5.08mm", "THT"),
    ("LOAD+ / LOAD-", "PhoenixContact_MSTBA_2,5_2-G-5,08_1x02_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757242", "MSTBA 2,5/ 2-G-5,08 header 2-pos 5.08mm", "THT"),
    ("WATER 1", "PhoenixContact_MSTBA_2,5_2-G-5,08_1x02_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757242", "MSTBA 2,5/ 2-G-5,08 header 2-pos 5.08mm", "THT"),
    ("WATER 2", "PhoenixContact_MSTBA_2,5_2-G-5,08_1x02_P5.08mm_Horizontal"):
        ("Phoenix Contact", "1757242", "MSTBA 2,5/ 2-G-5,08 header 2-pos 5.08mm", "THT"),
    ("GPS ANT (u.FL)", "U.FL_Hirose_U.FL-R-SMT-1_Vertical"): ("Hirose", "U.FL-R-SMT-1(10)", "u.FL receptacle", "SMD"),
    ("G5LE-1-E DC12", "Relay_SPDT_Omron-G5LE-1"): ("Omron", "G5LE-1-E DC12", "Relay SPDT 10A 12VDC coil", "THT"),
    ("10uH 6A", "L_Bourns_SRP1038C_10.0x10.0mm"): ("Bourns", "SRP1038C-100M", "Power inductor 10uH 10x10mm", "SMD"),
    ("27nH", "L_0603_1608Metric"): ("Murata", "LQG18HN27NJ00D", "RF inductor 27nH 0603", "SMD"),
    ("CSD18540Q5B", "TDSON-8-1"): ("Texas Instruments", "CSD18540Q5B", "N-MOSFET 60V 1.8mOhm SON5x6", "SMD"),
    ("AO3400A", "SOT-23"): ("Alpha & Omega", "AO3400A", "N-MOSFET 30V 5.7A SOT-23", "SMD"),
    ("BUK9M12-60EX", "LFPAK33"): ("Nexperia", "BUK9M12-60EX", "N-MOSFET 60V logic level LFPAK33", "SMD"),
    ("RESET", "SW_SPST_B3U-1000P"): ("Omron", "B3U-1000P", "Tactile switch SMD", "SMD"),
    ("BOOT", "SW_SPST_B3U-1000P"): ("Omron", "B3U-1000P", "Tactile switch SMD", "SMD"),
    ("LM74700QDBVRQ1", "SOT-23-6"): ("Texas Instruments", "LM74700QDBVRQ1", "Ideal diode controller SOT-23-6", "SMD"),
    ("TPS54360DDA", "TI_SO-PowerPAD-8_ThermalVias"): ("Texas Instruments", "TPS54360DDAR", "Buck converter 60V 3.5A HSOP-8", "SMD"),
    ("AMS1117-3.3", "SOT-223-3_TabPin2"): ("Advanced Monolithic Systems", "AMS1117-3.3", "LDO 3.3V 1A SOT-223", "SMD"),
    ("ESP32-S3-WROOM-1-N16", "ESP32-S3-WROOM-1"): ("Espressif", "ESP32-S3-WROOM-1-N16", "Wi-Fi/BLE module, 16MB flash, NO PSRAM", "SMD"),
    ("USBLC6-2SC6", "SOT-23-6"): ("STMicroelectronics", "USBLC6-2SC6", "USB ESD protection SOT-23-6", "SMD"),
    ("TJA1051T/3", "SOIC-8_3.9x4.9mm_P1.27mm"): ("NXP", "TJA1051T/3,118", "CAN transceiver SOIC-8", "SMD"),
    ("THVD1450D", "SOIC-8_3.9x4.9mm_P1.27mm"): ("Texas Instruments", "THVD1450DR", "RS-485 transceiver SOIC-8", "SMD"),
    ("MAX-M10S", "ublox_MAX"): ("u-blox", "MAX-M10S-00B", "GNSS module", "SMD"),
    ("LIS3DH", "LGA-16_3x3mm_P0.5mm_LayoutBorder3x5y"): ("STMicroelectronics", "LIS3DHTR", "Accelerometer LGA-16", "SMD"),
    ("ADS1115IDGS", "TSSOP-10_3x3mm_P0.5mm"): ("Texas Instruments", "ADS1115IDGSR", "16-bit ADC VSSOP-10", "SMD"),
    ("TCA9535PWR", "TSSOP-24_4.4x7.8mm_P0.65mm"): ("Texas Instruments", "TCA9535PWR", "I2C I/O expander TSSOP-24", "SMD"),
}
JUMPER = ("Any", "2.54mm 1x02 male header, gold", "Pin header 1x02 vertical (jumper position). Shunt NOT fitted", "THT")
RES_SERIES = ("Yageo", "RC0603FR-07{}L", "Resistor 0603 1% 0.1W", "SMD")

# Items shipped loose with the boards (not soldered)
LOOSE = [
    ("Phoenix Contact", "1757022", "MSTB 2,5/ 3-ST-5,08 plug, 3-pos (mates with 1757255)", 8),
    ("Phoenix Contact", "1757019", "MSTB 2,5/ 2-ST-5,08 plug, 2-pos (mates with 1757242)", 8),
    ("Any", "2.54mm jumper shunt", "Jumper shunt for JP1-JP10", 10),
]


def res_code(v):
    """Yageo RC0603 resistance code: 10k -> 10K, 4.7k -> 4K7, 100 -> 100R, 60.4 -> 60R4."""
    v = v.strip()
    if v.lower().endswith("k"):
        n = v[:-1]
        return n.replace(".", "K") if "." in n else n + "K"
    return v.replace(".", "R") if "." in v else v + "R"


def run(*a):
    subprocess.run(a, check=True, capture_output=True)


def main():
    if os.path.exists(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)

    # 1. board-relative place origin, then pick-and-place
    script = (
        "import pcbnew\n"
        f"b = pcbnew.LoadBoard(r'{PCB}')\n"
        f"b.GetDesignSettings().SetAuxOrigin(pcbnew.VECTOR2I(pcbnew.FromMM({BOARD_ORIGIN[0]}), pcbnew.FromMM({BOARD_ORIGIN[1]})))\n"
        f"pcbnew.SaveBoard(r'{PCB}', b)\n"
    )
    run(KIPY, "-c", script)
    raw_pos = os.path.join(PRJ, "build", "pos_board_origin.csv")
    run(CLI, "pcb", "export", "pos", "--format", "csv", "--units", "mm", "--side", "both",
        "--use-drill-file-origin", "-o", raw_pos, PCB)
    with open(raw_pos, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    placed_refs = []
    with open(os.path.join(OUT, f"03_PickPlace_VanController_{REV}.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Designator", "Mid X(mm)", "Mid Y(mm)", "Layer", "Rotation", "Value", "Package"])
        for r in rows:
            if r["Ref"] in DNP or r["Ref"].startswith("H"):
                continue
            if SELF_THT and r["Package"].startswith(THT_PKG):
                continue
            placed_refs.append(r["Ref"])
            w.writerow([r["Ref"], f"{float(r['PosX']):.3f}", f"{float(r['PosY']):.3f}",
                        "Top" if r["Side"] == "top" else "Bottom", f"{float(r['Rot']):.1f}", r["Val"], r["Package"]])
    placed = len(placed_refs)

    # 2. BOM
    raw_bom = os.path.join(PRJ, "build", "bom_raw.csv")
    run(CLI, "sch", "export", "bom", "--fields", "Reference,Value,Footprint,${QUANTITY}",
        "--labels", "Refs,Value,Footprint,Qty", "--group-by", "Value,Footprint", "-o", raw_bom, SCH)
    with open(raw_bom, encoding="utf-8") as f:
        groups = list(csv.DictReader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    head = ["Item", "Designator", "Qty per board", "Manufacturer", "Manufacturer Part Number", "Description",
            "Value", "Package / Footprint", "Mount type", "Fit", "Notes"]
    ws.append(head)
    unmatched = []
    item = 0
    smd = tht = 0
    for g in groups:
        refs, val, fp = g["Refs"], g["Value"], g["Footprint"].split(":")[-1]
        qty = int(g["Qty"])
        if refs.startswith("H"):
            continue
        if fp.startswith("R_0603"):
            m, pn, d, t = RES_SERIES
            pn = pn.format(res_code(val))
            d = f"{d} {val} ohm"
        elif fp.startswith("PinHeader_1x02"):
            m, pn, d, t = JUMPER
        elif (val, fp) in PARTS:
            m, pn, d, t = PARTS[(val, fp)]
        elif refs.split(",")[0] in DNP:
            m, pn, d, t = "-", "-", "M5 stud pad (no component)", "-"
        else:
            m = pn = d = t = ""
            unmatched.append((val, fp))
        first_ref = refs.split(",")[0].split("-")[0]
        fit = "DNP" if first_ref in DNP else "Fit"
        if SELF_THT and t == "THT":
            fit = "Ship loose"
        note = DNP.get(first_ref, "")
        if "verify" in pn:
            note = "Footprint is Littelfuse FLR 178.6165 (ATO, 30 A). Please quote an equivalent PCB ATO holder if needed."
        if fp.startswith(("R_", "C_", "CP_")) and m in ("Yageo", "Samsung", "Murata", "TDK", "Panasonic"):
            note = (note + " " if note else "") + "Equivalent to the same spec acceptable."
        item += 1
        ws.append([item, refs, qty, m, pn, d, val, fp, t, fit, note])
        if fit == "Fit":
            smd += qty
        elif fit == "Ship loose":
            tht += qty
    ws.append([])
    ws.append(["", "SHIP LOOSE (not soldered)"])
    for m, pn, d, q in LOOSE:
        item += 1
        ws.append([item, "-", q, m, pn, d, "", "", "Loose", "Ship loose", "Per board"])
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, wdt in enumerate([6, 34, 8, 22, 34, 46, 18, 44, 10, 10, 60], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(os.path.join(OUT, f"02_BOM_VanController_{REV}.xlsx"))
    if unmatched:
        print("UNMATCHED BOM lines:", unmatched)

    # 3. fabrication + documentation files
    gz = os.path.join(OUT, f"01_Gerbers_VanController_{REV}.zip")
    with zipfile.ZipFile(gz, "w", zipfile.ZIP_DEFLATED) as z:
        gdir = os.path.join(PRJ, "fab", "gerbers")
        for fn in sorted(os.listdir(gdir)):
            z.write(os.path.join(gdir, fn), fn)
    shutil.copy(os.path.join(PRJ, "docs", "VanController_assembly_top.pdf"),
                os.path.join(OUT, f"04_Assembly_Drawing_Top_{REV}.pdf"))
    shutil.copy(os.path.join(PRJ, "docs", "VanController_schematic.pdf"), os.path.join(OUT, f"05_Schematic_{REV}.pdf"))
    shutil.copy(os.path.join(PRJ, "docs", "3d_top.png"), os.path.join(OUT, f"06_3D_Top_{REV}.png"))
    shutil.copy(os.path.join(PRJ, "docs", "3d_iso.png"), os.path.join(OUT, f"06_3D_Iso_{REV}.png"))
    shutil.copy(os.path.join(PRJ, "fab", "VanController.step"), os.path.join(OUT, f"07_3D_Model_{REV}.step"))

    enc = os.path.join(OUT, "08_Enclosure_3D_print_optional")
    os.makedirs(enc)
    for fn in ("VanController_base.stl", "VanController_lid.stl"):
        shutil.copy(os.path.join(PRJ, "enclosure", fn), enc)
    with open(os.path.join(enc, "PRINT_NOTES.txt"), "w", encoding="utf-8") as f:
        f.write(PRINT_NOTES)
    with open(os.path.join(OUT, "00_READ_ME_ORDER_NOTES.txt"), "w", encoding="utf-8") as f:
        f.write(NOTES.format(rev=REV, smd=smd, tht=tht, placed=placed))

    z = shutil.make_archive(OUT, "zip", os.path.dirname(OUT), os.path.basename(OUT))
    print(f"package: {OUT}\nzip: {z}\nSMD parts/board: {smd}, THT parts/board: {tht}, placements: {placed}")


NOTES = """VanController {rev} - PCB fabrication + turnkey assembly order
=============================================================

Files
-----
01_Gerbers_*.zip          Gerber RS-274X + Excellon drill (KiCad 10). Layers: F.Cu, In1.Cu, In2.Cu, B.Cu,
                          mask, paste, silkscreen (top/bottom), Edge.Cuts, drill map.
02_BOM_*.xlsx             Bill of materials with manufacturer part numbers, fit/DNP and loose items.
03_PickPlace_*.csv        Centroid file, mm, origin = bottom-left corner of the board, KiCad rotation convention.
04_Assembly_Drawing_*.pdf Top-side assembly drawing with reference designators.
05_Schematic_*.pdf        Full schematic.
06_3D_*.png, 07_*.step    3D views / model for reference.
08_Enclosure_*            OPTIONAL separate 3D-printing order: enclosure base + lid STL and print notes.

PCB fabrication
---------------
Size:                230.0 x 135.0 mm, rectangular
Layers:              4 (stack-up: F.Cu signal / In1.Cu GND plane / In2.Cu +12V plane / B.Cu signal)
Material / Tg:       FR-4, Tg 150 or higher
Thickness:           1.6 mm
Outer copper:        1 oz (35 um). (2 oz is not usable: the design has 0.15 mm spacing at fine-pitch ICs)
Inner copper:        1 oz (35 um)
Min track/space:     0.15 / 0.15 mm
Min hole:            0.2 mm (ESP32 module thermal vias), standard vias 0.3 mm
Surface finish:      ENIG
Solder mask:         Green (any colour acceptable), white silkscreen
Via treatment:       Tented
Impedance control:   Not required
Quantity:            [your choice, e.g. 5 boards]

Assembly
--------
Type:                Turnkey (please source all parts), top side only
Parts per board:     {smd} SMD placements ({placed} centroid entries). Through-hole: 0 soldered.
Quantity:            [your choice, e.g. 2 or 3 boards assembled]
Through-hole parts:  DO NOT SOLDER. Please source them (BOM, Fit = "Ship loose") and ship them loose with
                     the boards: relays K1-K6, terminal headers J5-J18, J21-J22, fuse holders F1-F13,
                     pin headers J4, J20, JP1-JP10 ({tht} parts per board). Customer solders them.
DNP (do not fit):    J1, J2 (M5 stud pads - bolts fitted by customer), H1-H4 (mounting holes).
Do not fit fuses:    fuse holders only; no blade fuses in F1-F13.
Do not fit shunts:   JP1-JP10 jumper shunts shipped loose, not fitted.
Ship loose:          8 x Phoenix 1757022, 8 x Phoenix 1757019 plugs, 10 x jumper shunts per board.

Polarity and orientation
------------------------
- LEDs D4, D5, D12-D17, D40; diodes D1-D3, D6-D11, D24-D29: orientation per assembly drawing/silkscreen.
  D1 (SMDJ26CA) is bidirectional.
- C2 is a polarised aluminium electrolytic: + marking per silkscreen.
- U4 ESP32-S3-WROOM-1-N16: antenna end overhangs/faces the right board edge; keep the antenna area free.
  Variant must be N16 (no PSRAM) - do NOT substitute R8/R16V octal-PSRAM variants.
- U9 LIS3DH (LGA-16 3x3 mm, hidden joints): X-ray inspection requested. U10/U11: AOI is sufficient.
- Fine-pitch ICs have 0.15 mm pad gaps: no solder-mask dams needed; open the mask across those pad areas.

Testing
-------
No functional test required; please report any DFM issues before production.

Substitutions
-------------
Resistors and ceramic capacitors: any equivalent to the stated value, tolerance, voltage and dielectric.
All other parts: please ask before substituting.
"""

PRINT_NOTES = """Separate 3D printing order (optional).
Material: MJF PA12 (nylon) preferred, or PETG/ASA FDM. Not PLA (van interior temperatures).
Quantity: 1 x base + 1 x lid per assembled board. Units: mm. Outside 251 x 142 x 35 mm.
Hardware (customer fits): 8 x M3 heat-set inserts (4.0 mm hole), 4 x M3x8 screws for the lid,
4 x M3x6 screws for the PCB, 2 x M16 cable glands, 1 x SMA bulkhead to u.FL pigtail.
"""

if __name__ == "__main__":
    main()
